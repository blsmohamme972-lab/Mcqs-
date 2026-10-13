"""Extract every slide of a PowerPoint into per-slide folders.

Usage:
  python3 extract_slides.py deck.pptx [--out slides/<deck-name>]

For each slide NN it writes slides/<deck>/slide_NN/:
  content.md   title, text (with bullet levels), tables, speaker notes
  slide.png    the slide rendered as an image (needs LibreOffice + pdftoppm)
  media/       pictures embedded in the slide
and slides/<deck>/slides.json with the same text for every slide.
"""
import argparse
import json
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def slug(s):
    return re.sub(r"[^A-Za-z0-9]+", "-", s).strip("-").lower() or "deck"


def shape_text(shape, lines, media_dir, n_media):
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for s in shape.shapes:
            n_media = shape_text(s, lines, media_dir, n_media)
        return n_media
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        n_media += 1
        media_dir.mkdir(parents=True, exist_ok=True)
        (media_dir / f"image_{n_media}.{shape.image.ext}").write_bytes(shape.image.blob)
        return n_media
    if shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            t = "".join(r.text for r in p.runs).strip()
            if t:
                lines.append("  " * p.level + "- " + t)
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            lines.append("| " + " | ".join(c.text.strip() for c in row.cells) + " |")
    return n_media


def render(pptx, out_dir, count):
    """Render slides to PNG via LibreOffice -> PDF -> pdftoppm. Skips if tools are missing."""
    if not (shutil.which("soffice") and shutil.which("pdftoppm")):
        return False
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", tmp, str(pptx)],
                       check=True, capture_output=True, timeout=600)
        pdf = next(Path(tmp).glob("*.pdf"))
        subprocess.run(["pdftoppm", "-png", "-r", "80", str(pdf), f"{tmp}/p"], check=True, timeout=600)
        pages = sorted(Path(tmp).glob("p-*.png"), key=lambda p: int(p.stem.split("-")[-1]))
        for i, page in enumerate(pages[:count], 1):
            shutil.move(page, out_dir / f"slide_{i:02d}" / "slide.png")
    return True


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx", type=Path)
    ap.add_argument("--out", type=Path)
    a = ap.parse_args()

    out = a.out or Path(__file__).parent / "slides" / slug(a.pptx.stem)
    out.mkdir(parents=True, exist_ok=True)
    prs = Presentation(a.pptx)
    slides = []
    for n, slide in enumerate(prs.slides, 1):
        d = out / f"slide_{n:02d}"
        d.mkdir(exist_ok=True)
        title = slide.shapes.title.text.strip() if slide.shapes.title is not None else ""
        lines, n_media = [], 0
        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            n_media = shape_text(shape, lines, d / "media", n_media)
        notes = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
        md = [f"# Slide {n:02d}: {title or '(no title)'}", "", *lines]
        if notes:
            md += ["", "## Speaker notes", "", notes]
        (d / "content.md").write_text("\n".join(md) + "\n")
        slides.append({"number": n, "title": title, "text": lines, "notes": notes, "images": n_media})

    (out / "slides.json").write_text(json.dumps(slides, indent=2, ensure_ascii=False))
    rendered = render(a.pptx.resolve(), out, len(slides))
    print(f"{len(slides)} slides -> {out} (rendered: {rendered})")


if __name__ == "__main__":
    main()
